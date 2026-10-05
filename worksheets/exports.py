import base64
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.template.loader import render_to_string
from django.templatetags.static import static

from projects.models import Project
from .schema import PHASES


def document_context(worksheet, audience, *, preview=False):
    logo = Path(settings.BASE_DIR, "static/icons/brand-print.svg").read_bytes()
    exercises = [{**item, "number": index, "phase_label": dict(PHASES)[item["phase"]],
                  "lines": range(item["answer_lines"])} for index, item in enumerate(worksheet.payload["exercises"], 1)]
    return {"worksheet": worksheet, "payload": worksheet.payload, "source": worksheet.source,
            "brief": worksheet.brief, "teacher": audience == "teacher", "preview": preview,
            "language": dict(Project.Language.choices).get(worksheet.source["language"], worksheet.source["language"]),
            "level": worksheet.brief.get("level") or worksheet.source["level"],
            "pages": [exercises[index:index + 3] for index in range(0, len(exercises), 3)],
            "print_css": Path(settings.BASE_DIR, "static/css/worksheet-print.css").read_text(encoding="utf-8"),
            "logo": static("icons/brand-print.svg") if preview else "data:image/svg+xml;base64," + base64.b64encode(logo).decode()}


def export_pdf(worksheet, audience):
    from weasyprint import HTML
    from weasyprint.urls import URLFetcher, URLFetcherResponse
    context = document_context(worksheet, audience)
    class LocalLogoOnly(URLFetcher):
        def fetch(self, url, headers=None):
            if url != context["logo"]:
                raise ValueError("Externe Ressourcen sind im Arbeitsblatt nicht erlaubt.")
            return URLFetcherResponse(url, base64.b64decode(url.split(",", 1)[1]), {"Content-Type": "image/svg+xml"})
    return HTML(string=render_to_string("worksheets/document.html", context),
                url_fetcher=LocalLogoOnly(allowed_protocols=["data"], fail_on_errors=True)).write_pdf()


def export_docx(worksheet, audience):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.4), Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    section.header_distance, section.footer_distance = Cm(.8), Cm(.9)
    for name, size in (("Normal", 11), ("Title", 23), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Arial", Pt(size), RGBColor(17, 17, 17)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.18
        borders = style.element.get_or_add_pPr().find(qn("w:pBdr"))
        if borders is not None:
            borders.getparent().remove(borders)
    header = section.header.paragraphs[0]
    header.add_run().add_picture(str(Path(settings.BASE_DIR, "static/icons/brand-print.png")), width=Cm(.85))
    header.add_run("  Sprachplattform").bold = True
    header.add_run("  ·  Hörtexte · Hörspiele · Unterrichtsmaterial").font.size = Pt(8)
    context = document_context(worksheet, audience)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(("Lehrkräftefassung" if context["teacher"] else "Schülerfassung") + " · Seite ").font.size = Pt(8)
    field = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, end):
        field._r.append(element)

    def paragraph(text="", style=None, *, language=None):
        result = document.add_paragraph(text, style)
        if (language or worksheet.brief.get("instruction_language")) == "ar":
            bidi = OxmlElement("w:bidi")
            result._p.get_or_add_pPr().append(bidi)
        return result

    for index, tasks in enumerate(context["pages"]):
        if index:
            document.add_page_break()
        paragraph(worksheet.payload["title"], "Title")
        paragraph(f'{context["language"]} · {context["level"]} · {worksheet.brief["target_group"]}')
        if not context["teacher"]:
            paragraph("Name: _______________________   Klasse: ______   Datum: __________")
        if index == 0:
            paragraph("Lernziel: " + worksheet.payload["learning_goal"])
            paragraph(worksheet.payload["introduction"])
        for task in tasks:
            paragraph(f'{task["number"]}. {task["phase_label"]}', "Heading 2").paragraph_format.keep_with_next = True
            paragraph(task["question"]).paragraph_format.keep_with_next = True
            if context["teacher"]:
                paragraph("Lösung / Beispielantwort: " + task["answer"])
                if task["evidence"]:
                    paragraph("Textbeleg: „" + task["evidence"] + "“", language=worksheet.source["language"])
                if task["source_segments"]:
                    paragraph("Sprechbeiträge: " + ", ".join(map(str, task["source_segments"])))
            else:
                for option in task["options"]:
                    paragraph("☐  " + option)
                if not task["options"]:
                    for _ in task["lines"]:
                        line = paragraph()
                        line.paragraph_format.space_after = Pt(0)
                        line.paragraph_format.space_before = Pt(0)
                        line.paragraph_format.line_spacing = Pt(22)
                        borders = OxmlElement("w:pBdr")
                        bottom = OxmlElement("w:bottom")
                        for key, value in (("val", "single"), ("sz", "3"), ("color", "CCCCCC")):
                            bottom.set(qn("w:" + key), value)
                        borders.append(bottom)
                        between = OxmlElement("w:between")
                        for key, value in (("val", "single"), ("sz", "3"), ("color", "CCCCCC")):
                            between.set(qn("w:" + key), value)
                        borders.append(between)
                        line._p.get_or_add_pPr().append(borders)
    if context["teacher"]:
        document.add_page_break()
        paragraph("Hörtext · Textgrundlage", "Title")
        paragraph(worksheet.source["title"])
        for index, segment in enumerate(worksheet.source["segments"], 1):
            paragraph(f'{index}. {segment["speaker"]}: {segment["text"]}', language=worksheet.source["language"])
    output = BytesIO()
    document.save(output)
    return output.getvalue()
